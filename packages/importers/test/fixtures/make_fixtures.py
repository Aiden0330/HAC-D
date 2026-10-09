#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate the synthetic Excel fixtures for the HAC-D availability-spreadsheet importer.

Run from the repository root:

    uv run --with openpyxl python packages/importers/test/fixtures/make_fixtures.py

Outputs (all next to this script):

    series_basic_en.xlsx          F1   series system, English headers, Qty as exponent
    parallel_switchover_zh.xlsx   F2   FIT -> lambda, 1+1 with switchover success, Chinese headers
    kofn_combin.xlsx              F3   2-of-3 / 3-of-4 in explicit, COMBIN and BINOM.DIST form
    mttr_auto_manual.xlsx         F4   MTTR mixture p*T_auto + (1-p)*T_manual
    markov_1plus1.xlsx            F5   1+1 hot standby solved with MMULT / MINVERSE / TRANSPOSE
    hierarchy_levels_zh.xlsx      F6   tree given by a 层级 column (0, 1, 1.1, 1.1.1, ...)
    hierarchy_indent_en.xlsx      F7   the same tree given by indentation + row outline levels
    inconsistent_rare_event.xlsx  F8   1 - sum(U) next to PRODUCT(A), visibly different
    percent_units.xlsx            F9   availability in %, AFR in %/year
    manifest.json                      layout, expected tree, expected method tags, expected values

Everything here is SYNTHETIC: round-number MTBF / MTTR / FIT values chosen to exercise the
importer.  Nothing is vendor, paper or measured data.

Formulas are written as strings beginning with "=" and openpyxl stores them without cached
values (Excel, LibreOffice or the importer's engine computes them).  Every formula cell has an
expected value computed below with plain IEEE-754 doubles following the formula's own
operation order, so this script doubles as the reference evaluator for the fixture tests.

The generated files are byte-for-byte reproducible for a given openpyxl version: document
timestamps are pinned and the zip container is re-packed with fixed entry timestamps.
"""
from __future__ import annotations

import io
import json
import math
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.styles import Alignment, Font
from openpyxl.utils import FORMULAE, get_column_letter
from openpyxl.worksheet.properties import Outline

HERE = Path(__file__).resolve().parent
FIXED_DATE = "2026-10-09"
FIXED_DT = datetime(2026, 10, 9, 0, 0, 0)
BOLD = Font(bold=True)
TITLE_FONT = Font(bold=True, size=12)

# Method tags exactly as the importer tests expect them.
METHOD_TAGS = (
    "availability_mtbf_mttr",      # A = MTBF/(MTBF+MTTR)
    "availability_lambda_mttr",    # A = 1/(1+λ·MTTR)
    "unavailability_lambda_mttr",  # U = λ·MTTR
    "fit_to_lambda",               # λ = FIT·1e-9
    "afr_to_mtbf",                 # MTBF = 8760/AFR
    "series_product",              # PRODUCT(A…) or chained multiplication
    "series_rare_event_sum",       # 1 − SUM(U…)
    "parallel_product_complement", # 1 − PRODUCT(1 − A_i) or 1 − (1 − A)^n
    "kofn_explicit_sum",
    "kofn_combin",
    "kofn_binom_dist",
    "standby_exp",
    "switchover_weighted",         # c·A_sw + (1 − c)·A_nosw
    "mttr_mixture_auto_manual",    # p·T_auto + (1 − p)·T_manual
    "markov_mmult_minverse",
    "percent_to_fraction",
)

# Functions newer than the Excel 2007 specification must be stored with the "_xlfn." prefix.
# (SheetJS strips the prefix when reading unless {xlfn: true} is passed.)
POST_2007_FUNCTIONS = {"BINOM.DIST"}


# --------------------------------------------------------------------------------------
# Excel arithmetic in plain Python floats (same IEEE-754 doubles Excel uses)
# --------------------------------------------------------------------------------------

def xl_product(values):
    r = 1.0
    for v in values:
        r = r * v
    return r


def xl_sum(values):
    r = 0.0
    for v in values:
        r = r + v
    return r


def xl_combin(n, k):
    return float(math.comb(n, k))


def xl_binom_dist_cumulative(k, n, p):
    """BINOM.DIST(k, n, p, TRUE) = P(X <= k)."""
    return xl_sum(xl_combin(n, i) * p ** i * (1 - p) ** (n - i) for i in range(k + 1))


def xl_transpose(m):
    return [list(col) for col in zip(*m)]


def xl_mmult(a, b):
    return [[xl_sum(a[i][t] * b[t][j] for t in range(len(b))) for j in range(len(b[0]))]
            for i in range(len(a))]


def xl_minverse(m):
    """Gauss-Jordan elimination with partial pivoting (Excel's MINVERSE uses an LU solve;
    results agree to a few ulps for a 3x3 matrix)."""
    n = len(m)
    a = [list(map(float, row)) + [1.0 if i == j else 0.0 for j in range(n)] for i, row in enumerate(m)]
    for col in range(n):
        piv = max(range(col, n), key=lambda r: abs(a[r][col]))
        a[col], a[piv] = a[piv], a[col]
        p = a[col][col]
        if p == 0.0:
            raise ZeroDivisionError("singular matrix")
        a[col] = [x / p for x in a[col]]
        for r in range(n):
            if r != col and a[r][col] != 0.0:
                f = a[r][col]
                a[r] = [x - f * y for x, y in zip(a[r], a[col])]
    return [row[n:] for row in a]


# --------------------------------------------------------------------------------------
# Fixture recorder
# --------------------------------------------------------------------------------------

def ref(ws, addr):
    return f"{ws.title}!{addr}"


class Fixture:
    def __init__(self, fid, filename, title, language):
        self.id = fid
        self.file = filename
        self.title = title
        self.language = language
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.formulas = {}   # "Sheet!A1" -> "=..."
        self.values = {}     # "Sheet!A1" -> float
        self.inputs = {}     # "Sheet!A1" -> every non-formula value written (checked on read-back)
        self.entry = {}      # manifest entry (filled by the builder)

    def sheet(self, name):
        return self.wb.create_sheet(name)

    def put(self, ws, addr, value, *, bold=False, fmt=None, align=None):
        c = ws[addr]
        c.value = value
        if bold:
            c.font = BOLD
        if fmt:
            c.number_format = fmt
        if align:
            c.alignment = align
        if isinstance(value, str) and value.startswith("="):
            key = ref(ws, addr)
            assert key not in self.formulas, key
            self.formulas[key] = value
        elif value is not None:
            self.inputs[ref(ws, addr)] = value
        return c

    def expect(self, ws, addr, value):
        key = ref(ws, addr)
        assert key in self.formulas, f"{key} is not a formula cell"
        assert key not in self.values, key
        self.values[key] = float(value)

    def header(self, ws, row, labels):
        for i, label in enumerate(labels, start=1):
            self.put(ws, f"{get_column_letter(i)}{row}", label, bold=True)

    def widths(self, ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def finish(self):
        missing = set(self.formulas) - set(self.values)
        extra = set(self.values) - set(self.formulas)
        assert not missing, f"{self.id}: formulas without expected value: {sorted(missing)}"
        assert not extra, f"{self.id}: expected values without formula: {sorted(extra)}"
        for f in self.formulas.values():
            check_formula_syntax(f)


def check_formula_syntax(formula):
    """Static sanity check: balanced parentheses and every function name known to the
    Excel 2007 specification (or a post-2007 function carrying the _xlfn. prefix)."""
    tok = Tokenizer(formula)
    depth = 0
    for t in tok.items:
        if t.type == Token.FUNC:
            if t.subtype == Token.OPEN:
                depth += 1
                name = t.value[:-1].upper()
                if name.startswith("_XLFN."):
                    assert name[6:] in POST_2007_FUNCTIONS, f"unexpected _xlfn function {name} in {formula}"
                else:
                    assert name in FORMULAE, f"unknown function {name} in {formula}"
                    assert name not in POST_2007_FUNCTIONS, f"{name} needs the _xlfn. prefix in {formula}"
            else:
                depth -= 1
        elif t.type == Token.PAREN:
            depth += 1 if t.subtype == Token.OPEN else -1
    assert depth == 0, f"unbalanced parentheses in {formula}"


def node(name, kind, source_row, *, count=1, qty_cell=None, k=None, n=None, params=None,
         derived=None, cells=None, children=None, **extra):
    d = {
        "name": name,
        "kind": kind,
        "count": count,
        "qty_cell": qty_cell,
        "k": k,
        "n": n,
        "params": params or {},
        "derived": derived or {},
        "cells": cells or {},
        "source_row": source_row,
    }
    d.update(extra)
    d["children"] = children or []
    return d


def method(tag, cells, note=None, optional_cells=None):
    assert tag in METHOD_TAGS, tag
    d = {"tag": tag, "cells": list(cells)}
    if optional_cells:
        d["optional_cells"] = list(optional_cells)
    if note:
        d["note"] = note
    return d


def a_mtbf(mtbf, mttr):
    return mtbf / (mtbf + mttr)


# --------------------------------------------------------------------------------------
# F1  series_basic_en.xlsx
# --------------------------------------------------------------------------------------

def build_f1():
    fx = Fixture("F1", "series_basic_en.xlsx",
                 "Series system, English headers, per-unit availability raised to the quantity", "en")
    ws = fx.sheet("Series")
    fx.header(ws, 1, ["Component", "Qty", "MTBF (h)", "MTTR (h)", "Availability", "Unavailability"])
    fx.widths(ws, [30, 6, 12, 10, 16, 16])
    rows = [
        ("GPU node", 8, 50000, 6),
        ("Leaf switch", 4, 200000, 4),
        ("Spine switch", 2, 300000, 4),
        ("Storage array", 1, 100000, 12),
        ("Power feed", 1, 500000, 8),
        ("Cooling loop", 1, 150000, 24),
    ]
    leaves = []
    a_vals, u_vals = [], []
    for r, (name, qty, mtbf, mttr) in enumerate(rows, start=2):
        fx.put(ws, f"A{r}", name)
        fx.put(ws, f"B{r}", qty)
        fx.put(ws, f"C{r}", mtbf)
        fx.put(ws, f"D{r}", mttr)
        fx.put(ws, f"E{r}", f"=(C{r}/(C{r}+D{r}))^B{r}")
        fx.put(ws, f"F{r}", f"=1-E{r}")
        a = a_mtbf(mtbf, mttr) ** qty
        fx.expect(ws, f"E{r}", a)
        fx.expect(ws, f"F{r}", 1 - a)
        a_vals.append(a)
        u_vals.append(1 - a)
        leaves.append(node(name, "leaf", r, count=qty, qty_cell=qty,
                           params={"mtbf_h": mtbf, "mttr_h": mttr},
                           cells={"name": ref(ws, f"A{r}"), "count": ref(ws, f"B{r}"),
                                  "mtbf_h": ref(ws, f"C{r}"), "mttr_h": ref(ws, f"D{r}"),
                                  "availability": ref(ws, f"E{r}"), "unavailability": ref(ws, f"F{r}")}))
    fx.put(ws, "A8", "System (series)", bold=True)
    fx.put(ws, "E8", "=PRODUCT(E2:E7)")
    fx.put(ws, "F8", "=1-E8")
    a_sys = xl_product(a_vals)
    fx.expect(ws, "E8", a_sys)
    fx.expect(ws, "F8", 1 - a_sys)
    fx.put(ws, "A9", "System (rare-event approx.)", bold=True)
    fx.put(ws, "E9", "=1-SUM(F2:F7)")
    fx.put(ws, "F9", "=1-E9")
    a_rare = 1 - xl_sum(u_vals)
    fx.expect(ws, "E9", a_rare)
    fx.expect(ws, "F9", 1 - a_rare)

    tree = node("System (series)", "series", 8,
                cells={"availability": ref(ws, "E8"), "unavailability": ref(ws, "F8")},
                alt_cells={"series_rare_event_sum": ref(ws, "E9")},
                children=leaves)
    fx.entry = {
        "sheets": [ws.title], "sheet": ws.title, "header_row": 1, "data_rows": [2, 9],
        "headers": {get_column_letter(i): ws.cell(1, i).value for i in range(1, 7)},
        "columns": {"name": "A", "count": "B", "mtbf_h": "C", "mttr_h": "D",
                    "availability": "E", "unavailability": "F"},
        "merged_cells": [], "outline_levels": {},
        "expected_tree": tree,
        "expected_methods": [
            method("availability_mtbf_mttr", [ref(ws, f"E{r}") for r in range(2, 8)],
                   "per-unit A = C/(C+D) raised to Qty: (C/(C+D))^B; the exponent is the series multiplicity"),
            method("series_product", [ref(ws, "E8")]),
            method("series_rare_event_sum", [ref(ws, "E9")]),
        ],
        "untagged_formula_cells": [
            {"pattern": "1-<availability cell>", "cells": [ref(ws, f"F{r}") for r in range(2, 10)],
             "why": "complement U = 1 - A; no catalogue tag"},
        ],
        "tolerance": {"rel": 1e-12, "abs": 1e-15},
        "ambiguities": [
            "Qty is applied as an exponent inside the Availability formula ((C/(C+D))^B), so the leaf parameters are per unit and count = Qty.",
            "Two system rows: row 8 (exact PRODUCT) and row 9 (1 - SUM(U) rare-event approximation) describe the same system; the tree root is row 8 and row 9 is an alternative-method cell.",
        ],
        "extra": {"rare_event_vs_exact_abs_difference": a_sys - a_rare},
    }
    fx.finish()
    return fx


# --------------------------------------------------------------------------------------
# F2  parallel_switchover_zh.xlsx
# --------------------------------------------------------------------------------------

def build_f2():
    fx = Fixture("F2", "parallel_switchover_zh.xlsx",
                 "FIT to lambda, 1+1 redundancy with switchover success, Chinese headers", "zh")
    ws = fx.sheet("可用度")
    fx.header(ws, 1, ["部件", "数量", "失效率(FIT)", "失效率(/h)", "MTTR(h)", "单机可用度",
                      "冗余方式", "倒换成功率", "组可用度"])
    fx.widths(ws, [22, 6, 12, 14, 9, 16, 10, 12, 18])
    rows = [
        # name, qty, FIT, MTTR, redundancy, switchover success
        ("参数面交换机", 2, 5000, 4, "1+1", 0.98),
        ("管理服务器", 2, 8000, 8, "1+1", 0.95),
        ("存储控制器", 1, 20000, 6, "无", None),
    ]
    children = []
    group_vals = []
    for r, (name, qty, fit, mttr, red, c) in enumerate(rows, start=2):
        fx.put(ws, f"A{r}", name)
        fx.put(ws, f"B{r}", qty)
        fx.put(ws, f"C{r}", fit)
        fx.put(ws, f"D{r}", f"=C{r}*1E-9")
        fx.put(ws, f"E{r}", mttr)
        fx.put(ws, f"F{r}", f"=1/(1+D{r}*E{r})")
        fx.put(ws, f"G{r}", red)
        lam = fit * 1e-9
        a_unit = 1 / (1 + lam * mttr)
        fx.expect(ws, f"D{r}", lam)
        fx.expect(ws, f"F{r}", a_unit)
        cells = {"name": ref(ws, f"A{r}"), "count": ref(ws, f"B{r}"), "fit": ref(ws, f"C{r}"),
                 "lambda_per_h": ref(ws, f"D{r}"), "mttr_h": ref(ws, f"E{r}"),
                 "unit_availability": ref(ws, f"F{r}"), "redundancy": ref(ws, f"G{r}"),
                 "availability": ref(ws, f"I{r}")}
        if red == "1+1":
            fx.put(ws, f"H{r}", c)
            fx.put(ws, f"I{r}", f"=H{r}*(1-(1-F{r})^2)+(1-H{r})*F{r}")
            a_group = c * (1 - (1 - a_unit) ** 2) + (1 - c) * a_unit
            cells["switchover_success"] = ref(ws, f"H{r}")
            children.append(node(name, "parallel", r, count=1, qty_cell=qty, k=1, n=2, homogeneous=True,
                                 params={"fit": fit, "mttr_h": mttr, "switchover_success": c, "redundancy": red},
                                 derived={"lambda_per_h": lam, "unit_availability": a_unit},
                                 cells=cells))
        else:
            fx.put(ws, f"I{r}", f"=F{r}")
            a_group = a_unit
            children.append(node(name, "leaf", r, count=1, qty_cell=qty,
                                 params={"fit": fit, "mttr_h": mttr, "redundancy": red},
                                 derived={"lambda_per_h": lam, "unit_availability": a_unit},
                                 cells=cells))
        fx.expect(ws, f"I{r}", a_group)
        group_vals.append(a_group)
    fx.put(ws, "A5", "系统可用度", bold=True)
    fx.put(ws, "I5", "=PRODUCT(I2:I4)")
    fx.expect(ws, "I5", xl_product(group_vals))

    tree = node("系统可用度", "series", 5, cells={"availability": ref(ws, "I5")}, children=children)
    fx.entry = {
        "sheets": [ws.title], "sheet": ws.title, "header_row": 1, "data_rows": [2, 5],
        "headers": {get_column_letter(i): ws.cell(1, i).value for i in range(1, 10)},
        "columns": {"name": "A", "count": "B", "fit": "C", "lambda_per_h": "D", "mttr_h": "E",
                    "unit_availability": "F", "redundancy": "G", "switchover_success": "H",
                    "availability": "I"},
        "merged_cells": [], "outline_levels": {},
        "expected_tree": tree,
        "expected_methods": [
            method("fit_to_lambda", [ref(ws, f"D{r}") for r in (2, 3, 4)], "λ = FIT * 1E-9"),
            method("availability_lambda_mttr", [ref(ws, f"F{r}") for r in (2, 3, 4)]),
            method("parallel_product_complement", [ref(ws, "I2"), ref(ws, "I3")],
                   "sub-expression 1-(1-F)^2 inside the switchover weighting; the exponent 2 gives n"),
            method("switchover_weighted", [ref(ws, "I2"), ref(ws, "I3")],
                   "H*(A_parallel) + (1-H)*(A_unit) with H = 倒换成功率"),
            method("series_product", [ref(ws, "I5")]),
        ],
        "untagged_formula_cells": [
            {"pattern": "=<unit availability cell>", "cells": [ref(ws, "I4")],
             "why": "no redundancy (冗余方式 = 无): the group availability is a pass-through of the unit availability"},
        ],
        "tolerance": {"rel": 1e-12, "abs": 1e-15},
        "ambiguities": [
            "数量 = 2 on a 1+1 row counts the two units of the pair (n = 2), not two groups in series; the exponent 2 in the group formula confirms it. The tree therefore has count = 1, n = 2 for those rows.",
            "倒换成功率 is blank on the row without redundancy.",
            "Two availability-like columns: F (单机可用度, per unit) and I (组可用度, what the parent consumes). The tree's availability cell is column I.",
        ],
        "extra": {},
    }
    fx.finish()
    return fx


# --------------------------------------------------------------------------------------
# F3  kofn_combin.xlsx
# --------------------------------------------------------------------------------------

def build_f3():
    fx = Fixture("F3", "kofn_combin.xlsx",
                 "2-of-3 and 3-of-4 groups in explicit-sum, COMBIN and BINOM.DIST form", "en")
    ws = fx.sheet("kofn")
    fx.header(ws, 1, ["Component", "Qty (n)", "MTBF (h)", "MTTR (h)", "Unit availability",
                      "Redundancy (k/n)", "Group availability (explicit sum)",
                      "Group availability (COMBIN)", "Group availability (BINOM.DIST)"])
    fx.widths(ws, [22, 8, 10, 10, 16, 14, 30, 28, 30])
    rows = [
        # name, n, k, MTBF, MTTR
        ("Power shelf PSU", 3, 2, 20000, 48),
        ("CDU pump", 4, 3, 30000, 72),
        ("Management switch", 1, 1, 200000, 4),
    ]
    children = []
    g_vals, h_vals, i_vals = [], [], []
    for r, (name, n, k, mtbf, mttr) in enumerate(rows, start=2):
        fx.put(ws, f"A{r}", name)
        fx.put(ws, f"B{r}", n)
        fx.put(ws, f"C{r}", mtbf)
        fx.put(ws, f"D{r}", mttr)
        fx.put(ws, f"E{r}", f"=C{r}/(C{r}+D{r})")
        fx.put(ws, f"F{r}", f"{k}/{n}")
        a = a_mtbf(mtbf, mttr)
        fx.expect(ws, f"E{r}", a)
        cells = {"name": ref(ws, f"A{r}"), "count": ref(ws, f"B{r}"), "mtbf_h": ref(ws, f"C{r}"),
                 "mttr_h": ref(ws, f"D{r}"), "unit_availability": ref(ws, f"E{r}"),
                 "redundancy": ref(ws, f"F{r}"), "availability": ref(ws, f"G{r}")}
        if n == 3 and k == 2:
            fx.put(ws, f"G{r}", f"=E{r}^3+3*E{r}^2*(1-E{r})")
            fx.put(ws, f"H{r}", f"=COMBIN(3,3)*E{r}^3*(1-E{r})^0+COMBIN(3,2)*E{r}^2*(1-E{r})^1")
            fx.put(ws, f"I{r}", f"=1-_xlfn.BINOM.DIST(1,3,E{r},TRUE)")
            g = a ** 3 + 3 * a ** 2 * (1 - a)
            h = xl_combin(3, 3) * a ** 3 * (1 - a) ** 0 + xl_combin(3, 2) * a ** 2 * (1 - a) ** 1
            i = 1 - xl_binom_dist_cumulative(1, 3, a)
        elif n == 4 and k == 3:
            fx.put(ws, f"G{r}", f"=E{r}^4+4*E{r}^3*(1-E{r})")
            fx.put(ws, f"H{r}", f"=COMBIN(4,4)*E{r}^4*(1-E{r})^0+COMBIN(4,3)*E{r}^3*(1-E{r})^1")
            fx.put(ws, f"I{r}", f"=1-_xlfn.BINOM.DIST(2,4,E{r},TRUE)")
            g = a ** 4 + 4 * a ** 3 * (1 - a)
            h = xl_combin(4, 4) * a ** 4 * (1 - a) ** 0 + xl_combin(4, 3) * a ** 3 * (1 - a) ** 1
            i = 1 - xl_binom_dist_cumulative(2, 4, a)
        else:
            fx.put(ws, f"G{r}", f"=E{r}")
            fx.put(ws, f"H{r}", f"=E{r}")
            fx.put(ws, f"I{r}", f"=E{r}")
            g = h = i = a
        fx.expect(ws, f"G{r}", g)
        fx.expect(ws, f"H{r}", h)
        fx.expect(ws, f"I{r}", i)
        g_vals.append(g)
        h_vals.append(h)
        i_vals.append(i)
        if n > 1:
            children.append(node(name, "kofn", r, count=1, qty_cell=n, k=k, n=n, homogeneous=True,
                                 params={"mtbf_h": mtbf, "mttr_h": mttr, "redundancy": f"{k}/{n}"},
                                 derived={"unit_availability": a},
                                 cells=dict(cells, availability_alt=[ref(ws, f"H{r}"), ref(ws, f"I{r}")])))
        else:
            children.append(node(name, "leaf", r, count=1, qty_cell=n,
                                 params={"mtbf_h": mtbf, "mttr_h": mttr, "redundancy": f"{k}/{n}"},
                                 derived={"unit_availability": a},
                                 cells=dict(cells, availability_alt=[ref(ws, f"H{r}"), ref(ws, f"I{r}")])))
    fx.put(ws, "A5", "System (series)", bold=True)
    fx.put(ws, "G5", "=PRODUCT(G2:G4)")
    fx.put(ws, "H5", "=PRODUCT(H2:H4)")
    fx.put(ws, "I5", "=PRODUCT(I2:I4)")
    fx.expect(ws, "G5", xl_product(g_vals))
    fx.expect(ws, "H5", xl_product(h_vals))
    fx.expect(ws, "I5", xl_product(i_vals))

    tree = node("System (series)", "series", 5,
                cells={"availability": ref(ws, "G5"), "availability_alt": [ref(ws, "H5"), ref(ws, "I5")]},
                children=children)
    fx.entry = {
        "sheets": [ws.title], "sheet": ws.title, "header_row": 1, "data_rows": [2, 5],
        "headers": {get_column_letter(i): ws.cell(1, i).value for i in range(1, 10)},
        "columns": {"name": "A", "count": "B", "mtbf_h": "C", "mttr_h": "D", "unit_availability": "E",
                    "redundancy": "F", "availability": "G"},
        "merged_cells": [], "outline_levels": {},
        "expected_tree": tree,
        "expected_methods": [
            method("availability_mtbf_mttr", [ref(ws, f"E{r}") for r in (2, 3, 4)]),
            method("kofn_explicit_sum", [ref(ws, "G2"), ref(ws, "G3")],
                   "A^n + n*A^(n-1)*(1-A): the terms i = n and i = n-1 of the k-of-n sum"),
            method("kofn_combin", [ref(ws, "H2"), ref(ws, "H3")]),
            method("kofn_binom_dist", [ref(ws, "I2"), ref(ws, "I3")],
                   "1-BINOM.DIST(k-1, n, A, TRUE) = P(at least k of n up); stored as _xlfn.BINOM.DIST"),
            method("series_product", [ref(ws, "G5"), ref(ws, "H5"), ref(ws, "I5")]),
        ],
        "untagged_formula_cells": [
            {"pattern": "=<unit availability cell>", "cells": [ref(ws, "G4"), ref(ws, "H4"), ref(ws, "I4")],
             "why": "1/1 row: group availability is a pass-through of the unit availability"},
        ],
        "tolerance": {"rel": 1e-12, "abs": 1e-15},
        "ambiguities": [
            "Columns G, H and I compute the same group availability three ways (they agree to ~1e-16); the tree's availability cell is G, H and I are listed under availability_alt. An importer should pick one and report the others as cross-checks, not create three groups.",
            "Qty (n) holds the number of units in the voting group; k and n also appear as text 'k/n' in column F. count = 1 in the tree.",
            "BINOM.DIST is stored as _xlfn.BINOM.DIST in the file (required for post-2007 functions); SheetJS strips the prefix by default.",
            "Three system cells (G5, H5, I5) are all PRODUCT of their own column.",
        ],
        "extra": {"availability_variants": ["G", "H", "I"]},
    }
    fx.finish()
    return fx


# --------------------------------------------------------------------------------------
# F4  mttr_auto_manual.xlsx
# --------------------------------------------------------------------------------------

def build_f4():
    fx = Fixture("F4", "mttr_auto_manual.xlsx",
                 "Effective MTTR as a mixture of automatic and manual recovery", "zh/en")
    ws = fx.sheet("MTTR")
    fx.header(ws, 1, ["Component", "Qty", "λ (/h)", "自动恢复比例 p", "MTTR自动(h)", "MTTR人工(h)",
                      "MTTR有效(h)", "可用度", "备注"])
    fx.widths(ws, [26, 6, 12, 14, 12, 12, 12, 16, 44])
    rows = [
        # name, qty, aggregate lambda (/h) for the row, p_auto, T_auto, T_manual, note
        ("GPU 节点 (XID/HBM 故障)", 32, 6.4e-4, 0.7, 0.25, 6,
         "λ 为本行 32 台合计 / aggregate λ of the 32 nodes in this row"),
        ("参数面交换机", 2, 1.0e-5, 0.9, 0.1, 4,
         "λ 为本行 2 台合计 / aggregate λ of the 2 switches"),
        ("并行文件系统", 4, 4.0e-5, 0.5, 0.5, 12,
         "λ 为本行 4 台合计 / aggregate λ of the 4 storage nodes"),
        ("调度/管理面", 1, 5.0e-6, 0.8, 0.05, 8, None),
    ]
    children = []
    a_vals = []
    for r, (name, qty, lam, p, t_auto, t_man, note) in enumerate(rows, start=2):
        fx.put(ws, f"A{r}", name)
        fx.put(ws, f"B{r}", qty)
        fx.put(ws, f"C{r}", lam)
        fx.put(ws, f"D{r}", p)
        fx.put(ws, f"E{r}", t_auto)
        fx.put(ws, f"F{r}", t_man)
        fx.put(ws, f"G{r}", f"=D{r}*E{r}+(1-D{r})*F{r}")
        fx.put(ws, f"H{r}", f"=1/(1+C{r}*G{r})")
        if note:
            fx.put(ws, f"I{r}", note)
        t_eff = p * t_auto + (1 - p) * t_man
        a = 1 / (1 + lam * t_eff)
        fx.expect(ws, f"G{r}", t_eff)
        fx.expect(ws, f"H{r}", a)
        a_vals.append(a)
        children.append(node(name, "leaf", r, count=1, qty_cell=qty, lambda_scope="row_aggregate",
                             params={"lambda_per_h": lam, "auto_fraction": p, "mttr_auto_h": t_auto,
                                     "mttr_manual_h": t_man},
                             derived={"mttr_effective_h": t_eff},
                             cells={"name": ref(ws, f"A{r}"), "count": ref(ws, f"B{r}"),
                                    "lambda_per_h": ref(ws, f"C{r}"), "auto_fraction": ref(ws, f"D{r}"),
                                    "mttr_auto_h": ref(ws, f"E{r}"), "mttr_manual_h": ref(ws, f"F{r}"),
                                    "mttr_effective_h": ref(ws, f"G{r}"), "availability": ref(ws, f"H{r}"),
                                    **({"notes": ref(ws, f"I{r}")} if note else {})}))
    fx.put(ws, "A6", "系统 (System)", bold=True)
    fx.put(ws, "H6", "=PRODUCT(H2:H5)")
    fx.expect(ws, "H6", xl_product(a_vals))

    tree = node("系统 (System)", "series", 6, cells={"availability": ref(ws, "H6")}, children=children)
    fx.entry = {
        "sheets": [ws.title], "sheet": ws.title, "header_row": 1, "data_rows": [2, 6],
        "headers": {get_column_letter(i): ws.cell(1, i).value for i in range(1, 10)},
        "columns": {"name": "A", "count": "B", "lambda_per_h": "C", "auto_fraction": "D",
                    "mttr_auto_h": "E", "mttr_manual_h": "F", "mttr_effective_h": "G",
                    "availability": "H", "notes": "I"},
        "merged_cells": [], "outline_levels": {},
        "expected_tree": tree,
        "expected_methods": [
            method("mttr_mixture_auto_manual", [ref(ws, f"G{r}") for r in range(2, 6)],
                   "D*E + (1-D)*F with D = 自动恢复比例"),
            method("availability_lambda_mttr", [ref(ws, f"H{r}") for r in range(2, 6)],
                   "1/(1+λ·MTTR_eff) using the mixture MTTR"),
            method("series_product", [ref(ws, "H6")]),
        ],
        "untagged_formula_cells": [],
        "tolerance": {"rel": 1e-12, "abs": 1e-15},
        "ambiguities": [
            "Qty (column B) is not referenced by any formula: the λ column is the aggregate failure rate of the whole row (see the 备注 column), so the tree uses count = 1 and lambda_scope = row_aggregate. Contrast F1, where Qty is an exponent.",
            "Mixed zh/en headers on one row.",
        ],
        "extra": {},
    }
    fx.finish()
    return fx


# --------------------------------------------------------------------------------------
# F5  markov_1plus1.xlsx
# --------------------------------------------------------------------------------------

def build_f5():
    fx = Fixture("F5", "markov_1plus1.xlsx",
                 "1+1 hot standby solved as a 3-state Markov chain with MMULT/MINVERSE/TRANSPOSE", "en")
    ws = fx.sheet("Markov")
    fx.widths(ws, [52, 16, 16, 16, 4, 8])
    lam, mttr = 2e-5, 4
    mu = 1 / mttr

    ws.merge_cells("A1:F1")
    fx.put(ws, "A1", "Head-node pair, 1+1 hot standby, 3-state Markov chain (synthetic parameters)", bold=True)
    ws["A1"].font = TITLE_FONT
    fx.put(ws, "A3", "λ (/h) per node")
    fx.put(ws, "B3", lam)
    fx.put(ws, "A4", "MTTR (h)")
    fx.put(ws, "B4", mttr)
    fx.put(ws, "A5", "μ (/h) = 1/MTTR")
    fx.put(ws, "B5", "=1/B4")
    fx.expect(ws, "B5", mu)

    # Generator Q, rows = from-state, columns = to-state, states [2 up, 1 up, 0 up]; one repair crew.
    fx.put(ws, "A7", "Q (from \\ to)", bold=True)
    for col, label in zip("BCD", ["2 up", "1 up", "0 up"]):
        fx.put(ws, f"{col}7", label, bold=True, align=Alignment(horizontal="center"))
    fx.put(ws, "A8", "2 up", bold=True)
    fx.put(ws, "B8", "=-2*$B$3")
    fx.put(ws, "C8", "=2*$B$3")
    fx.put(ws, "D8", 0)
    fx.put(ws, "A9", "1 up", bold=True)
    fx.put(ws, "B9", "=$B$5")
    fx.put(ws, "C9", "=-($B$3+$B$5)")
    fx.put(ws, "D9", "=$B$3")
    fx.put(ws, "A10", "0 up", bold=True)
    fx.put(ws, "B10", 0)
    fx.put(ws, "C10", "=$B$5")
    fx.put(ws, "D10", "=-$B$5")
    q = [[-2 * lam, 2 * lam, 0.0],
         [mu, -(lam + mu), lam],
         [0.0, mu, -mu]]
    fx.expect(ws, "B8", q[0][0])
    fx.expect(ws, "C8", q[0][1])
    fx.expect(ws, "B9", q[1][0])
    fx.expect(ws, "C9", q[1][1])
    fx.expect(ws, "D9", q[1][2])
    fx.expect(ws, "C10", q[2][1])
    fx.expect(ws, "D10", q[2][2])

    # M = transpose(Q) with the last equation replaced by the normalisation sum(pi) = 1; b = (0, 0, 1).
    fx.put(ws, "A12", "M = Qᵀ with the last row replaced by Σπ = 1", bold=True)
    fx.put(ws, "F12", "b", bold=True, align=Alignment(horizontal="center"))
    qt = xl_transpose(q)
    m = [qt[0], qt[1], [1.0, 1.0, 1.0]]
    for i, r in enumerate((13, 14)):
        fx.put(ws, f"A{r}", f"balance equation {i + 1}")
        for j, col in enumerate("BCD"):
            fx.put(ws, f"{col}{r}", f"=INDEX(TRANSPOSE($B$8:$D$10),{i + 1},{j + 1})")
            fx.expect(ws, f"{col}{r}", qt[i][j])
    fx.put(ws, "A15", "Σπ = 1")
    for col in "BCD":
        fx.put(ws, f"{col}15", 1)
    fx.put(ws, "F13", 0)
    fx.put(ws, "F14", 0)
    fx.put(ws, "F15", 1)
    b = [[0.0], [0.0], [1.0]]

    fx.put(ws, "A17", "Steady state π = M⁻¹ b", bold=True)
    pi = xl_mmult(xl_minverse(m), b)
    pi = [row[0] for row in pi]
    for i, (r, label) in enumerate(zip((18, 19, 20), ["π(2 up)", "π(1 up)", "π(0 up)"])):
        fx.put(ws, f"A{r}", label)
        fx.put(ws, f"B{r}", f"=INDEX(MMULT(MINVERSE($B$13:$D$15),$F$13:$F$15),{i + 1},1)")
        fx.expect(ws, f"B{r}", pi[i])
    fx.put(ws, "A21", "check Σπ")
    fx.put(ws, "B21", "=SUM(B18:B20)")
    fx.expect(ws, "B21", xl_sum(pi))
    fx.put(ws, "A22", "Availability (≥ 1 up) = π(2 up) + π(1 up)", bold=True)
    fx.put(ws, "B22", "=B18+B19")
    fx.expect(ws, "B22", pi[0] + pi[1])
    fx.put(ws, "A23", "Unavailability (0 up)")
    fx.put(ws, "B23", "=B20")
    fx.expect(ws, "B23", pi[2])

    # Closed-form cross-check of the matrix solve (balance equations): pi1 = 2rho*pi2, pi0 = 2rho^2*pi2.
    rho = lam / mu
    pi2_cf = 1 / (1 + 2 * rho + 2 * rho ** 2)
    pi_cf = [pi2_cf, 2 * rho * pi2_cf, 2 * rho ** 2 * pi2_cf]
    for x, y in zip(pi, pi_cf):
        assert abs(x - y) < 1e-12, (pi, pi_cf)

    fx.put(ws, "A25", "Cross-check: independent units (equivalent to two repair crews)", bold=True)
    fx.put(ws, "A26", "Unit availability 1/(1+λ·MTTR)")
    fx.put(ws, "B26", "=1/(1+B3*B4)")
    a_unit = 1 / (1 + lam * mttr)
    fx.expect(ws, "B26", a_unit)
    fx.put(ws, "A27", "1+1 parallel 1−(1−A)^2")
    fx.put(ws, "B27", "=1-(1-B26)^2")
    a_par = 1 - (1 - a_unit) ** 2
    fx.expect(ws, "B27", a_par)
    fx.put(ws, "A28", "Δ (parallel − Markov)")
    fx.put(ws, "B28", "=B27-B22")
    fx.expect(ws, "B28", a_par - (pi[0] + pi[1]))

    fx.put(ws, "A30", "Cold standby, no repair, mission time t (h) — reliability, not availability", bold=True)
    fx.put(ws, "B30", 8760)
    fx.put(ws, "A31", "R(t) = e^(−λt)·(1+λt)")
    fx.put(ws, "B31", "=EXP(-B3*B30)*(1+B3*B30)")
    t = 8760
    fx.expect(ws, "B31", math.exp(-lam * t) * (1 + lam * t))

    tree = node("Head-node pair (1+1 hot standby)", "parallel", 22, count=1, qty_cell=None, k=1, n=2,
                homogeneous=True,
                params={"lambda_per_h": lam, "mttr_h": mttr},
                derived={"mu_per_h": mu, "unit_availability": a_unit},
                cells={"lambda_per_h": ref(ws, "B3"), "mttr_h": ref(ws, "B4"), "mu_per_h": ref(ws, "B5"),
                       "availability": ref(ws, "B22"), "unavailability": ref(ws, "B23"),
                       "state_probabilities": [ref(ws, "B18"), ref(ws, "B19"), ref(ws, "B20")]},
                evaluation={
                    "method": "ctmc",
                    "states": ["2 up", "1 up", "0 up"],
                    "generator_cells": ref(ws, "B8:D10"),
                    "repair_policy": "single repair crew: '0 up' -> '1 up' at μ (not 2μ); this is why B22 differs from the independent-units value B27 by B28",
                    "closed_form": "π(2 up) = 1/(1+2ρ+2ρ²), π(1 up) = 2ρ·π(2 up), π(0 up) = 2ρ²·π(2 up), ρ = λ/μ",
                })
    fx.entry = {
        "sheets": [ws.title], "sheet": ws.title, "header_row": None, "layout": "block", "data_rows": [3, 31],
        "headers": {},
        "columns": {"name": "A"},
        "merged_cells": ["A1:F1"], "outline_levels": {},
        "expected_tree": tree,
        "expected_methods": [
            method("markov_mmult_minverse", [ref(ws, "B18"), ref(ws, "B19"), ref(ws, "B20")],
                   "INDEX(MMULT(MINVERSE(M), b), r, 1) per cell, no CSE arrays; M is built in B13:D15 with "
                   "INDEX(TRANSPOSE(Q), r, c) and a row of ones, b in F13:F15",
                   optional_cells=[ref(ws, f"{c}{r}") for r in (13, 14) for c in "BCD"] + [ref(ws, "B22")]),
            method("availability_lambda_mttr", [ref(ws, "B26")]),
            method("parallel_product_complement", [ref(ws, "B27")]),
            method("standby_exp", [ref(ws, "B31")],
                   "EXP(-λt)*(1+λt): cold-standby mission reliability without repair; a different 口径 (reliability over t, not availability) and NOT part of the tree"),
        ],
        "untagged_formula_cells": [
            {"pattern": "generator matrix entries from λ and μ", "cells": [ref(ws, c) for c in ("B8", "C8", "B9", "C9", "D9", "C10", "D10")],
             "why": "Q entries (−2λ, 2λ, μ, −(λ+μ), λ, μ, −μ)"},
            {"pattern": "μ = 1/MTTR", "cells": [ref(ws, "B5")], "why": "rate from MTTR; no catalogue tag"},
            {"pattern": "sum / passthrough / difference", "cells": [ref(ws, "B21"), ref(ws, "B22"), ref(ws, "B23"), ref(ws, "B28")],
             "why": "Σπ check, availability = π(2 up)+π(1 up), unavailability = π(0 up), Δ vs independent units"},
        ],
        "tolerance": {"rel": 1e-9, "abs": 1e-13,
                      "note": "π(0 up) ≈ 1.28e-8 is obtained from an O(1) matrix solve, so Excel and any other solver agree with it only to ~1e-16 absolute (~1e-8 relative); compare the small cells (B20, B23, B28) absolutely"},
        "ambiguities": [
            "Block layout, no header row and no component table: the model is a single 1+1 group whose parameters are label/value pairs (A3:B5).",
            "Availability is the sum of two state probabilities (B22 = B18+B19), not a formula over component availabilities; the importer should recognise the MMULT/MINVERSE block and treat the chain as one node solved by CTMC.",
            "The chain uses one repair crew (0 up -> 1 up at μ). An engine that models the pair as two independent repairable units (2μ) reproduces B27, not B22; the difference B28 = B27 - B22 ≈ +6.4e-9 (independent units are more available than the single-crew chain) is intended.",
            "B31 is a mission-reliability formula (EXP) sitting on the same sheet; it must be tagged standby_exp but kept out of the availability tree.",
        ],
        "extra": {"state_probabilities_closed_form": pi_cf},
    }
    fx.finish()
    return fx


# --------------------------------------------------------------------------------------
# Shared tree for F6 / F7
# --------------------------------------------------------------------------------------

# level, depth, zh name, en name, qty, MTBF, MTTR, kind, redundancy zh, redundancy en, combine form, pow form
TREE_ROWS = [
    ("0",     0, "系统",                 "System",                        None, None,   None, "series",   "串联",        "series",           "product_list",            False),
    ("1",     1, "计算子系统",           "Compute subsystem",             None, None,   None, "series",   "串联",        "series",           "chain",                   False),
    ("1.1",   2, "计算机柜",             "Compute rack",                  None, None,   None, "series",   "串联",        "series",           "product_range",           False),
    ("1.1.1", 3, "GPU 节点",             "GPU node",                      8,    50000,  6,    "leaf",     "无",          "none",             None,                      True),
    ("1.1.2", 3, "机柜 ToR 交换机",      "Rack ToR switch",               2,    200000, 4,    "leaf",     "无",          "none",             None,                      True),
    ("1.1.3", 3, "机柜 PDU",             "Rack PDU",                      1,    400000, 8,    "leaf",     "无",          "none",             None,                      True),
    ("1.2",   2, "管理节点",             "Management node",               1,    100000, 8,    "leaf",     "无",          "none",             None,                      True),
    ("2",     1, "网络子系统",           "Network subsystem",             None, None,   None, "parallel", "1+1 (双平面)", "1+1 (dual plane)", "parallel_pairwise",       False),
    ("2.1",   2, "网络平面 A",           "Network plane A",               1,    150000, 4,    "leaf",     "无",          "none",             None,                      False),
    ("2.2",   2, "网络平面 B",           "Network plane B",               1,    150000, 4,    "leaf",     "无",          "none",             None,                      False),
    ("3",     1, "供电子系统",           "Power subsystem",               None, None,   None, "parallel", "2N",          "2N",               "parallel_product_u_range", False),
    ("3.1",   2, "A 路供电 (市电+UPS)",  "Power path A (utility + UPS)",  1,    80000,  12,   "leaf",     "无",          "none",             None,                      False),
    ("3.2",   2, "B 路供电 (市电+UPS)",  "Power path B (utility + UPS)",  1,    80000,  12,   "leaf",     "无",          "none",             None,                      False),
]


def tree_children():
    """index -> list of child indices, from the depth sequence."""
    children = {i: [] for i in range(len(TREE_ROWS))}
    stack = []
    for i, row in enumerate(TREE_ROWS):
        depth = row[1]
        while stack and TREE_ROWS[stack[-1]][1] >= depth:
            stack.pop()
        if stack:
            children[stack[-1]].append(i)
        stack.append(i)
    return children


def build_hierarchy(fx, ws, *, first_row, col, lang, level_col=None, indent=False, outline=False):
    """Write the shared tree into ws.  col maps role -> column letter.  Returns (tree, methods, untagged)."""
    children = tree_children()
    row_of = {i: first_row + i for i in range(len(TREE_ROWS))}
    A, U, Q, MTBF, MTTR = col["availability"], col["unavailability"], col["count"], col["mtbf_h"], col["mttr_h"]
    name_col, red_col, notes_col = col["name"], col["redundancy"], col["notes"]
    a_val = {}
    methods = {"availability_mtbf_mttr": [], "series_product": [], "parallel_product_complement": []}
    complement_cells = []
    nodes = {}

    # leaves first (post-order evaluation), then groups
    def evaluate(i):
        level, depth, zh, en, qty, mtbf, mttr, kind, red_zh, red_en, form, pow_form = TREE_ROWS[i]
        r = row_of[i]
        if kind == "leaf":
            if pow_form:
                a = a_mtbf(mtbf, mttr) ** qty
            else:
                a = a_mtbf(mtbf, mttr)
        else:
            kids = children[i]
            for k in kids:
                evaluate(k)
            vals = [a_val[k] for k in kids]
            if form == "product_list" or form == "product_range":
                a = xl_product(vals)
            elif form == "chain":
                a = vals[0]
                for v in vals[1:]:
                    a = a * v
            elif form == "parallel_pairwise":
                a = 1 - (1 - vals[0]) * (1 - vals[1])
            elif form == "parallel_product_u_range":
                a = 1 - xl_product([1 - v for v in vals])
            else:
                raise ValueError(form)
        a_val[i] = a

    evaluate(0)

    for i, (level, depth, zh, en, qty, mtbf, mttr, kind, red_zh, red_en, form, pow_form) in enumerate(TREE_ROWS):
        r = row_of[i]
        name = zh if lang == "zh" else en
        red = red_zh if lang == "zh" else red_en
        raw_name = ("  " * depth + name) if indent else name
        if level_col:
            level_value = int(level) if "." not in level else level
            fx.put(ws, f"{level_col}{r}", level_value)
        fx.put(ws, f"{name_col}{r}", raw_name, bold=(kind != "leaf"))
        if outline and depth > 0:
            ws.row_dimensions[r].outline_level = depth
        if qty is not None:
            fx.put(ws, f"{Q}{r}", qty)
            fx.put(ws, f"{MTBF}{r}", mtbf)
            fx.put(ws, f"{MTTR}{r}", mttr)
        fx.put(ws, f"{red_col}{r}", red)
        kids = children[i]
        kid_rows = [row_of[k] for k in kids]
        if kind == "leaf":
            if pow_form:
                f = f"=({MTBF}{r}/({MTBF}{r}+{MTTR}{r}))^{Q}{r}"
            else:
                f = f"={MTBF}{r}/({MTBF}{r}+{MTTR}{r})"
            methods["availability_mtbf_mttr"].append(ref(ws, f"{A}{r}"))
        elif form == "product_list":
            f = "=PRODUCT(" + ",".join(f"{A}{kr}" for kr in kid_rows) + ")"
            methods["series_product"].append(ref(ws, f"{A}{r}"))
        elif form == "chain":
            f = "=" + "*".join(f"{A}{kr}" for kr in kid_rows)
            methods["series_product"].append(ref(ws, f"{A}{r}"))
        elif form == "product_range":
            f = f"=PRODUCT({A}{kid_rows[0]}:{A}{kid_rows[-1]})"
            methods["series_product"].append(ref(ws, f"{A}{r}"))
        elif form == "parallel_pairwise":
            f = f"=1-(1-{A}{kid_rows[0]})*(1-{A}{kid_rows[1]})"
            methods["parallel_product_complement"].append(ref(ws, f"{A}{r}"))
        elif form == "parallel_product_u_range":
            f = f"=1-PRODUCT({U}{kid_rows[0]}:{U}{kid_rows[-1]})"
            methods["parallel_product_complement"].append(ref(ws, f"{A}{r}"))
        fx.put(ws, f"{A}{r}", f)
        fx.expect(ws, f"{A}{r}", a_val[i])
        fx.put(ws, f"{U}{r}", f"=1-{A}{r}")
        fx.expect(ws, f"{U}{r}", 1 - a_val[i])
        complement_cells.append(ref(ws, f"{U}{r}"))
        notes = {
            "product_list": "串联：PRODUCT 引用三个子系统" if lang == "zh" else "series: PRODUCT over the three subsystems",
            "chain": "串联：连乘" if lang == "zh" else "series: chained multiplication",
            "product_range": "串联：PRODUCT 连续区间" if lang == "zh" else "series: PRODUCT over a range",
            "parallel_pairwise": "并联：1-(1-A)(1-B)" if lang == "zh" else "parallel: 1-(1-A)(1-B)",
            "parallel_product_u_range": "并联：1-PRODUCT(不可用度)" if lang == "zh" else "parallel: 1-PRODUCT(unavailabilities)",
        }
        if form:
            fx.put(ws, f"{notes_col}{r}", notes[form])
        cells = {"name": ref(ws, f"{name_col}{r}"), "availability": ref(ws, f"{A}{r}"),
                 "unavailability": ref(ws, f"{U}{r}"), "redundancy": ref(ws, f"{red_col}{r}")}
        if level_col:
            cells["level"] = ref(ws, f"{level_col}{r}")
        extra = {}
        if level_col:
            extra["level"] = level
            extra["level_cell_type"] = "n" if "." not in level else "s"
        if indent:
            extra["name_cell_raw"] = raw_name
            extra["indent_depth"] = depth
        if outline:
            extra["outline_level"] = depth
        if kind == "leaf":
            cells.update({"count": ref(ws, f"{Q}{r}"), "mtbf_h": ref(ws, f"{MTBF}{r}"), "mttr_h": ref(ws, f"{MTTR}{r}")})
            nodes[i] = node(name, "leaf", r, count=(qty if pow_form else 1), qty_cell=qty,
                            params={"mtbf_h": mtbf, "mttr_h": mttr, "redundancy": red},
                            cells=cells, **extra)
        elif kind == "series":
            nodes[i] = node(name, "series", r, params={"redundancy": red}, cells=cells,
                            combine_formula_form=form, **extra)
        else:
            nodes[i] = node(name, "parallel", r, k=1, n=len(kids), params={"redundancy": red}, cells=cells,
                            combine_formula_form=form, **extra)
    for i in range(len(TREE_ROWS) - 1, -1, -1):
        nodes[i]["children"] = [nodes[k] for k in children[i]]
    untagged = [{"pattern": "1-<availability cell>", "cells": complement_cells,
                 "why": "complement U = 1 - A; no catalogue tag (but F6/F7 parallel 2N formulas reference these cells)"}]
    return nodes[0], methods, untagged, row_of


def build_f6():
    fx = Fixture("F6", "hierarchy_levels_zh.xlsx", "Tree given by a 层级 (level) column, Chinese headers", "zh")
    ws = fx.sheet("拓扑")
    ws.merge_cells("A1:I1")
    fx.put(ws, "A1", "AI 集群可用度模型 — 层级列示例（合成数据，非实测）", bold=True)
    ws["A1"].font = TITLE_FONT
    fx.header(ws, 2, ["层级", "名称", "数量", "MTBF(h)", "MTTR(h)", "可用度", "不可用度", "冗余方式", "备注"])
    fx.widths(ws, [8, 26, 6, 10, 9, 16, 16, 14, 34])
    col = {"level": "A", "name": "B", "count": "C", "mtbf_h": "D", "mttr_h": "E", "availability": "F",
           "unavailability": "G", "redundancy": "H", "notes": "I"}
    tree, methods, untagged, row_of = build_hierarchy(fx, ws, first_row=3, col=col, lang="zh", level_col="A")
    fx.entry = {
        "sheets": [ws.title], "sheet": ws.title, "header_row": 2, "data_rows": [3, 15],
        "headers": {get_column_letter(i): ws.cell(2, i).value for i in range(1, 10)},
        "columns": col,
        "merged_cells": ["A1:I1"], "outline_levels": {},
        "expected_tree": tree,
        "expected_methods": [
            method("availability_mtbf_mttr", methods["availability_mtbf_mttr"],
                   "rows 6-9 use (D/(D+E))^C (Qty as exponent), rows 11, 12, 14, 15 use D/(D+E)"),
            method("series_product", methods["series_product"],
                   "F3 = PRODUCT(F4,F10,F13) non-contiguous list; F4 = F5*F9 chained; F5 = PRODUCT(F6:F8) range"),
            method("parallel_product_complement", methods["parallel_product_complement"],
                   "F10 = 1-(1-F11)*(1-F12); F13 = 1-PRODUCT(G14:G15) over the 不可用度 column"),
        ],
        "untagged_formula_cells": untagged,
        "tolerance": {"rel": 1e-12, "abs": 1e-15},
        "ambiguities": [
            "Row 1 is a merged title (A1:I1); the header is row 2.",
            "The root row has 层级 = 0, which is not a dotted prefix of '1', '2', '3': rows whose level has no dot are children of the single level-0 row. The formula F3 = PRODUCT(F4,F10,F13) confirms it.",
            "Level cells are mixed types: 0, 1, 2, 3 are stored as numbers, dotted levels ('1.1', '1.1.1') as text. Real sheets also store 1.1 as the number 1.1 (and then '1.10' collides with '1.1'); the fixture avoids that collision but an importer must coerce numeric levels to strings.",
            "F13 (供电子系统 2N) references the 不可用度 column G14:G15 rather than the 可用度 column, so the formula graph must map G -> 1-F to identify the children rows 14 and 15.",
            "Mixed depth: 1.2 管理节点 is a leaf directly under 计算子系统 while 1.1 计算机柜 is a group.",
            "Redundancy text '2N' and '1+1 (双平面)' both mean k=1 of n=2.",
        ],
        "extra": {"row_of_level": {TREE_ROWS[i][0]: r for i, r in row_of.items()}},
    }
    fx.finish()
    return fx


def build_f7():
    fx = Fixture("F7", "hierarchy_indent_en.xlsx",
                 "Same tree as F6, hierarchy given by leading spaces in the name and row outline levels", "en")
    ws = fx.sheet("Topology")
    fx.header(ws, 1, ["Component", "Qty", "MTBF (h)", "MTTR (h)", "Availability", "Unavailability",
                      "Redundancy", "Notes"])
    fx.widths(ws, [36, 6, 10, 9, 16, 16, 18, 40])
    col = {"name": "A", "count": "B", "mtbf_h": "C", "mttr_h": "D", "availability": "E",
           "unavailability": "F", "redundancy": "G", "notes": "H"}
    tree, methods, untagged, row_of = build_hierarchy(fx, ws, first_row=2, col=col, lang="en",
                                                      indent=True, outline=True)
    # Parents sit above their children: tell Excel the summary rows are above (default is below).
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False)
    ws.sheet_format.outlineLevelRow = max(r[1] for r in TREE_ROWS)
    outline_levels = {str(row_of[i]): TREE_ROWS[i][1] for i in range(len(TREE_ROWS)) if TREE_ROWS[i][1] > 0}
    fx.entry = {
        "sheets": [ws.title], "sheet": ws.title, "header_row": 1, "data_rows": [2, 14],
        "headers": {get_column_letter(i): ws.cell(1, i).value for i in range(1, 9)},
        "columns": col,
        "merged_cells": [],
        "outline_levels": outline_levels,
        "outline_summary_below": False,
        "indent": {"unit": "2 leading spaces per depth level in the Component cell", "max_depth": 3},
        "expected_tree": tree,
        "expected_methods": [
            method("availability_mtbf_mttr", methods["availability_mtbf_mttr"],
                   "rows 5-8 use (C/(C+D))^B (Qty as exponent), rows 10, 11, 13, 14 use C/(C+D)"),
            method("series_product", methods["series_product"],
                   "E2 = PRODUCT(E3,E9,E12) non-contiguous list; E3 = E4*E8 chained; E4 = PRODUCT(E5:E7) range"),
            method("parallel_product_complement", methods["parallel_product_complement"],
                   "E9 = 1-(1-E10)*(1-E11); E12 = 1-PRODUCT(F13:F14) over the Unavailability column"),
        ],
        "untagged_formula_cells": untagged,
        "tolerance": {"rel": 1e-12, "abs": 1e-15},
        "ambiguities": [
            "No level column: depth = (leading spaces in Component) / 2 = Excel row outline level. Both cues agree; the root row has no outline level (0).",
            "Excel's default outline convention is 'summary row below detail'; this sheet sets summaryBelow = false because parents sit above their children. SheetJS exposes row levels via ws['!rows'][i].level; check whether it exposes outlinePr before relying on it.",
            "Leading spaces are stored in the shared string (xml:space='preserve'); the manifest's node.name is trimmed, node.name_cell_raw keeps the spaces.",
            "E12 (Power subsystem 2N) references the Unavailability column F13:F14; map F -> 1-E to find the children.",
        ],
        "extra": {"row_of_level": {TREE_ROWS[i][0]: r for i, r in row_of.items()}},
    }
    fx.finish()
    return fx


# --------------------------------------------------------------------------------------
# F8  inconsistent_rare_event.xlsx
# --------------------------------------------------------------------------------------

def build_f8():
    fx = Fixture("F8", "inconsistent_rare_event.xlsx",
                 "Series system whose headline uses 1 - sum(U) while an adjacent cell uses PRODUCT(A)", "en")
    ws = fx.sheet("Series")
    fx.header(ws, 1, ["Component", "Qty", "MTBF (h)", "MTTR (h)", "λ (/h)", "Unavailability (λ·MTTR)",
                      "Availability (exact)"])
    fx.widths(ws, [44, 6, 10, 10, 12, 22, 20])
    rows = [
        ("GPU node pool (degraded-mode threshold)", 1, 2000, 48),
        ("Interconnect fabric", 1, 4000, 72),
        ("Parallel file system", 1, 3000, 96),
        ("Scheduler / control plane", 1, 5000, 120),
        ("Facility (power + cooling)", 1, 6000, 240),
    ]
    children = []
    u_approx, a_exact = [], []
    for r, (name, qty, mtbf, mttr) in enumerate(rows, start=2):
        fx.put(ws, f"A{r}", name)
        fx.put(ws, f"B{r}", qty)
        fx.put(ws, f"C{r}", mtbf)
        fx.put(ws, f"D{r}", mttr)
        fx.put(ws, f"E{r}", f"=1/C{r}")
        fx.put(ws, f"F{r}", f"=E{r}*D{r}")
        fx.put(ws, f"G{r}", f"=C{r}/(C{r}+D{r})")
        lam = 1 / mtbf
        u = lam * mttr
        a = a_mtbf(mtbf, mttr)
        fx.expect(ws, f"E{r}", lam)
        fx.expect(ws, f"F{r}", u)
        fx.expect(ws, f"G{r}", a)
        u_approx.append(u)
        a_exact.append(a)
        children.append(node(name, "leaf", r, count=1, qty_cell=qty,
                             params={"mtbf_h": mtbf, "mttr_h": mttr},
                             derived={"lambda_per_h": lam, "unavailability_lambda_mttr": u},
                             cells={"name": ref(ws, f"A{r}"), "count": ref(ws, f"B{r}"),
                                    "mtbf_h": ref(ws, f"C{r}"), "mttr_h": ref(ws, f"D{r}"),
                                    "lambda_per_h": ref(ws, f"E{r}"), "unavailability": ref(ws, f"F{r}"),
                                    "availability": ref(ws, f"G{r}")}))
    fx.put(ws, "A8", "System availability (rare-event, 1 − ΣU)", bold=True)
    fx.put(ws, "F8", "=SUM(F2:F6)")
    fx.put(ws, "G8", "=1-SUM(F2:F6)")
    sum_u = xl_sum(u_approx)
    fx.expect(ws, "F8", sum_u)
    fx.expect(ws, "G8", 1 - sum_u)
    fx.put(ws, "A9", "Exact (ΠA)", bold=True)
    fx.put(ws, "F9", "=1-G9")
    fx.put(ws, "G9", "=PRODUCT(G2:G6)")
    prod_a = xl_product(a_exact)
    fx.expect(ws, "G9", prod_a)
    fx.expect(ws, "F9", 1 - prod_a)
    fx.put(ws, "A10", "Difference (exact − rare-event)", bold=True)
    fx.put(ws, "G10", "=G9-G8")
    diff = prod_a - (1 - sum_u)
    fx.expect(ws, "G10", diff)
    assert abs(diff) >= 1e-3, diff

    tree = node("System", "series", 9,
                cells={"availability": ref(ws, "G9"), "unavailability": ref(ws, "F9")},
                alt_cells={"series_rare_event_sum": ref(ws, "G8"), "headline": ref(ws, "G8")},
                children=children)
    fx.entry = {
        "sheets": [ws.title], "sheet": ws.title, "header_row": 1, "data_rows": [2, 10],
        "headers": {get_column_letter(i): ws.cell(1, i).value for i in range(1, 8)},
        "columns": {"name": "A", "count": "B", "mtbf_h": "C", "mttr_h": "D", "lambda_per_h": "E",
                    "unavailability": "F", "availability": "G"},
        "merged_cells": [], "outline_levels": {},
        "expected_tree": tree,
        "expected_methods": [
            method("availability_mtbf_mttr", [ref(ws, f"G{r}") for r in range(2, 7)]),
            method("unavailability_lambda_mttr", [ref(ws, f"F{r}") for r in range(2, 7)],
                   "U = λ·MTTR first-order approximation (λ·MTTR is 0.018-0.04 here, so it is visibly off)"),
            method("series_rare_event_sum", [ref(ws, "G8")], "the sheet's headline number",
                   optional_cells=[ref(ws, "F8")]),
            method("series_product", [ref(ws, "G9")]),
        ],
        "untagged_formula_cells": [
            {"pattern": "1/<MTBF cell>", "cells": [ref(ws, f"E{r}") for r in range(2, 7)],
             "why": "λ = 1/MTBF; no tag in the catalogue (closest: fit_to_lambda / afr_to_mtbf are other conversions)"},
            {"pattern": "SUM(U range)", "cells": [ref(ws, "F8")], "why": "ΣU helper next to the 1-ΣU cell"},
            {"pattern": "1-<availability cell>", "cells": [ref(ws, "F9")], "why": "complement"},
            {"pattern": "difference", "cells": [ref(ws, "G10")], "why": "exact minus rare-event"},
        ],
        "tolerance": {"rel": 1e-12, "abs": 1e-15},
        "ambiguities": [
            "The sheet's headline 'System availability' (G8) is the rare-event approximation 1-ΣU built on first-order U = λ·MTTR; the exact PRODUCT sits one row below (G9). The tree root is G9; G8 is recorded under alt_cells and the reconciliation table must show the difference (G10).",
            "Two stacked approximations: U = λ·MTTR per row (instead of MTTR/(MTBF+MTTR)) and 1-ΣU for the series.",
            "MTBF values (2000-6000 h) are deliberately low so the approximations are visibly wrong (≥ 1e-3).",
        ],
        "extra": {
            "inconsistency": {
                "rare_event_cell": ref(ws, "G8"), "exact_cell": ref(ws, "G9"), "difference_cell": ref(ws, "G10"),
                "rare_event_value": 1 - sum_u, "exact_value": prod_a, "exact_minus_rare_event": diff,
                "sum_of_lambda_mttr": sum_u,
            }
        },
    }
    fx.finish()
    return fx


# --------------------------------------------------------------------------------------
# F9  percent_units.xlsx
# --------------------------------------------------------------------------------------

def build_f9():
    fx = Fixture("F9", "percent_units.xlsx", "Availability entered in percent, AFR in %/year", "zh")
    ws = fx.sheet("单位")
    fx.header(ws, 1, ["部件", "数量", "AFR (%/年)", "MTBF (h)", "MTTR (h)", "可用度(%)", "可用度 (小数)", "备注"])
    fx.widths(ws, [22, 6, 11, 12, 9, 12, 16, 48])
    children = []
    g_vals = []

    def afr_row(r, name, qty, afr_pct, mttr, note):
        fx.put(ws, f"A{r}", name)
        fx.put(ws, f"B{r}", qty)
        fx.put(ws, f"C{r}", afr_pct)
        fx.put(ws, f"D{r}", f"=8760/(C{r}/100)")
        fx.put(ws, f"E{r}", mttr)
        fx.put(ws, f"F{r}", f"=G{r}*100", fmt="0.0000")
        fx.put(ws, f"G{r}", f"=(D{r}/(D{r}+E{r}))^B{r}")
        fx.put(ws, f"H{r}", note)
        mtbf = 8760 / (afr_pct / 100)
        g = (mtbf / (mtbf + mttr)) ** qty
        fx.expect(ws, f"D{r}", mtbf)
        fx.expect(ws, f"G{r}", g)
        fx.expect(ws, f"F{r}", g * 100)
        g_vals.append(g)
        children.append(node(name, "leaf", r, count=qty, qty_cell=qty,
                             params={"afr_pct_per_year": afr_pct, "mttr_h": mttr},
                             derived={"mtbf_h": mtbf},
                             cells={"name": ref(ws, f"A{r}"), "count": ref(ws, f"B{r}"),
                                    "afr_pct_per_year": ref(ws, f"C{r}"), "mtbf_h": ref(ws, f"D{r}"),
                                    "mttr_h": ref(ws, f"E{r}"), "availability_pct": ref(ws, f"F{r}"),
                                    "availability": ref(ws, f"G{r}"), "notes": ref(ws, f"H{r}")}))

    def sla_row(r, name, pct_value, note, *, percent_format=False):
        fx.put(ws, f"A{r}", name)
        fx.put(ws, f"B{r}", 1)
        if percent_format:
            fx.put(ws, f"F{r}", pct_value, fmt="0.00%")     # stored as a fraction, displayed as 99.99%
            fx.put(ws, f"G{r}", f"=F{r}")
            g = pct_value
        else:
            fx.put(ws, f"F{r}", pct_value, fmt="0.00")      # stored as 99.95, a plain number
            fx.put(ws, f"G{r}", f"=F{r}/100")
            g = pct_value / 100
        fx.put(ws, f"H{r}", note)
        fx.expect(ws, f"G{r}", g)
        g_vals.append(g)
        children.append(node(name, "leaf", r, count=1, qty_cell=1,
                             params={"availability_pct": pct_value} if not percent_format else {"availability": pct_value},
                             derived={"availability": g},
                             cells={"name": ref(ws, f"A{r}"), "count": ref(ws, f"B{r}"),
                                    "availability_pct": ref(ws, f"F{r}"), "availability": ref(ws, f"G{r}"),
                                    "notes": ref(ws, f"H{r}")},
                             availability_cell_number_format=("0.00%" if percent_format else "0.00")))

    afr_row(2, "GPU 服务器", 64, 12, 8, "AFR 换算 MTBF = 8760/(AFR/100)；可用度按数量连乘")
    sla_row(3, "存储服务 (SLA)", 99.95, "可用度以百分数录入（数值单元格 99.95），使用前 /100")
    sla_row(4, "网络服务 (SLA)", 99.9, "同上")
    sla_row(5, "机房供电 (SLA)", 0.9999, "单元格为百分比格式（显示 99.99%），存储值已是小数，不再 /100", percent_format=True)
    afr_row(6, "参数面交换机", 2, 3, 4, "AFR 换算 MTBF")
    fx.put(ws, "A7", "系统可用度", bold=True)
    fx.put(ws, "F7", "=G7*100", fmt="0.0000")
    fx.put(ws, "G7", "=PRODUCT(G2:G6)")
    g_sys = xl_product(g_vals)
    fx.expect(ws, "G7", g_sys)
    fx.expect(ws, "F7", g_sys * 100)

    tree = node("系统可用度", "series", 7, cells={"availability": ref(ws, "G7"), "availability_pct": ref(ws, "F7")},
                children=children)
    fx.entry = {
        "sheets": [ws.title], "sheet": ws.title, "header_row": 1, "data_rows": [2, 7],
        "headers": {get_column_letter(i): ws.cell(1, i).value for i in range(1, 9)},
        "columns": {"name": "A", "count": "B", "afr_pct_per_year": "C", "mtbf_h": "D", "mttr_h": "E",
                    "availability_pct": "F", "availability": "G", "notes": "H"},
        "merged_cells": [], "outline_levels": {},
        "cell_number_formats": {ref(ws, "F3"): "0.00", ref(ws, "F4"): "0.00", ref(ws, "F5"): "0.00%",
                                ref(ws, "F2"): "0.0000", ref(ws, "F6"): "0.0000", ref(ws, "F7"): "0.0000"},
        "expected_tree": tree,
        "expected_methods": [
            method("afr_to_mtbf", [ref(ws, "D2"), ref(ws, "D6")], "MTBF = 8760/(AFR/100) with AFR in %/year"),
            method("percent_to_fraction", [ref(ws, "G3"), ref(ws, "G4")], "A = F/100 for percent entered as a number",
                   optional_cells=[ref(ws, "D2"), ref(ws, "D6")]),
            method("availability_mtbf_mttr", [ref(ws, "G2"), ref(ws, "G6")], "(D/(D+E))^B with Qty as exponent"),
            method("series_product", [ref(ws, "G7")]),
        ],
        "untagged_formula_cells": [
            {"pattern": "<availability cell>*100", "cells": [ref(ws, "F2"), ref(ws, "F6"), ref(ws, "F7")],
             "why": "fraction -> percent for display; the inverse of percent_to_fraction, no catalogue tag"},
            {"pattern": "=<percent-formatted cell>", "cells": [ref(ws, "G5")],
             "why": "F5 stores 0.9999 with number format 0.00% (displays 99.99%), so no /100 is applied"},
        ],
        "tolerance": {"rel": 1e-12, "abs": 1e-15},
        "ambiguities": [
            "Column F (可用度(%)) mixes inputs and formulas: 99.95 and 99.9 are plain numbers (format 0.00) that the sheet divides by 100; 0.9999 in F5 has the percent number format and is used without /100; F2, F6, F7 are formulas (=G*100) for display.",
            "The percent/fraction decision must come from the formula (or the number format), not from the '%' in the header alone.",
            "D2 and D6 are formula cells (MTBF derived from AFR), so the leaf's mtbf_h is a computed value with provenance to the AFR cell.",
            "The /100 inside 8760/(C/100) may or may not be reported as percent_to_fraction; D2 and D6 are listed as optional_cells for that tag.",
            "Roles afr_pct_per_year and availability_pct are extensions beyond the base role list.",
        ],
        "extra": {},
    }
    fx.finish()
    return fx


# --------------------------------------------------------------------------------------
# Deterministic save and read-back verification
# --------------------------------------------------------------------------------------

def save_deterministic(wb, path):
    wb.properties.creator = "HAC-D make_fixtures.py"
    wb.properties.created = FIXED_DT
    wb.properties.modified = FIXED_DT   # openpyxl overwrites this at save time; fixed again below
    wb.save(path)
    with zipfile.ZipFile(path) as zin:
        entries = [(zi, zin.read(zi.filename)) for zi in zin.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for zi, data in entries:
            if zi.filename == "docProps/core.xml":
                data = re.sub(rb"<dcterms:modified([^>]*)>[^<]*</dcterms:modified>",
                              rb"<dcterms:modified\1>" + FIXED_DATE.encode() + rb"T00:00:00Z</dcterms:modified>",
                              data)
            nzi = zipfile.ZipInfo(zi.filename, date_time=(1980, 1, 1, 0, 0, 0))
            nzi.compress_type = zipfile.ZIP_DEFLATED
            nzi.external_attr = zi.external_attr
            zout.writestr(nzi, data)
    path.write_bytes(buf.getvalue())


def verify(fx, path):
    """Read the workbook back (formulas, not cached values) and check everything the manifest claims."""
    wb = load_workbook(path)          # data_only=False: formulas as strings
    assert wb.sheetnames == fx.entry["sheets"], (wb.sheetnames, fx.entry["sheets"])
    found = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    found[ref(ws, c.coordinate)] = c.value
    assert found == fx.formulas, (fx.id, set(found) ^ set(fx.formulas))
    # Every input value must read back exactly as written: openpyxl serialises floats with
    # "%.16g", which does not round-trip every double, and strings must keep leading spaces.
    for key, written in fx.inputs.items():
        sheet, addr = key.split("!")
        got = wb[sheet][addr].value
        assert got == written and type(got) is type(written), (fx.id, key, written, got)
    ws = wb[fx.entry["sheet"]]
    merged = sorted(str(m) for m in ws.merged_cells.ranges)
    assert merged == sorted(fx.entry["merged_cells"]), (merged, fx.entry["merged_cells"])
    levels = {str(r): d.outlineLevel for r, d in ws.row_dimensions.items() if d.outlineLevel}
    assert levels == fx.entry["outline_levels"], (levels, fx.entry["outline_levels"])
    if fx.entry.get("header_row"):
        hr = fx.entry["header_row"]
        for col, label in fx.entry["headers"].items():
            assert ws[f"{col}{hr}"].value == label, (col, label)
    # every cell the tree points at exists and holds something
    def walk(n):
        for role, cell in n["cells"].items():
            for c in (cell if isinstance(cell, list) else [cell]):
                sheet, addr = c.split("!")
                if ":" in addr:
                    continue
                assert wb[sheet][addr].value is not None, (fx.id, role, c)
        for ch in n["children"]:
            walk(ch)
    walk(fx.entry["expected_tree"])
    for m in fx.entry["expected_methods"]:
        for c in m["cells"] + m.get("optional_cells", []):
            assert c in fx.formulas, (fx.id, m["tag"], c)
    with zipfile.ZipFile(path) as z:
        xml = z.read(f"xl/worksheets/sheet1.xml")
        if fx.entry["outline_levels"]:
            assert b'outlineLevel="' in xml, "outline levels not written"
            assert b'summaryBelow="0"' in xml, "summaryBelow not written"
        if any("_xlfn." in f for f in fx.formulas.values()):
            assert b"_xlfn.BINOM.DIST" in xml
    return ws.max_row


def tree_rows(n, depth=0, out=None):
    out = [] if out is None else out
    tag = n["kind"]
    if n["kind"] in ("parallel", "kofn"):
        tag += f" {n['k']}/{n['n']}"
    out.append(f"{'  ' * depth}{n['name']} [{tag}, row {n['source_row']}, count {n['count']}]")
    for ch in n["children"]:
        tree_rows(ch, depth + 1, out)
    return out


def main():
    builders = [build_f1, build_f2, build_f3, build_f4, build_f5, build_f6, build_f7, build_f8, build_f9]
    fixtures = [b() for b in builders]
    workbooks = []
    coverage = {tag: [] for tag in METHOD_TAGS}
    print(f"Writing fixtures to {HERE}")
    for fx in fixtures:
        path = HERE / fx.file
        save_deterministic(fx.wb, path)
        max_row = verify(fx, path)
        for m in fx.entry["expected_methods"]:
            coverage[m["tag"]].append(fx.id)
        entry = {"id": fx.id, "file": fx.file, "title": fx.title, "language": fx.language}
        entry.update(fx.entry)
        entry["formulas"] = dict(sorted(fx.formulas.items()))
        entry["expected_values"] = dict(sorted(fx.values.items()))
        workbooks.append(entry)
        tags = ", ".join(m["tag"] for m in fx.entry["expected_methods"])
        root = fx.entry["expected_tree"]
        print(f"\n{fx.id}  {fx.file}  sheet={fx.entry['sheet']!r}  rows={max_row}  "
              f"formulas={len(fx.formulas)}  size={path.stat().st_size} B")
        print(f"    methods: {tags}")
        print(f"    root availability {root['cells']['availability']} = {fx.values[root['cells']['availability']]!r}")
        for line in tree_rows(root):
            print("    | " + line)

    manifest = {
        "version": "1.0",
        "generated": FIXED_DATE,
        "generator": "packages/importers/test/fixtures/make_fixtures.py",
        "command": "uv run --with openpyxl python packages/importers/test/fixtures/make_fixtures.py",
        "data_disclaimer": "All parameter values are synthetic round numbers chosen to exercise the importer; none are vendor, paper or measured data.",
        "conventions": {
            "cell_address": "'Sheet!A1' (sheet name without quotes, 1-based rows)",
            "formulas": "exact strings stored in the file; openpyxl writes no cached values, so a reader must evaluate them (post-2007 functions carry the _xlfn. prefix, which SheetJS strips by default)",
            "expected_values": "one number per formula cell, computed in Python with IEEE-754 doubles following the formula's own operation order; compare with the per-workbook tolerance (rel OR abs)",
            "expected_methods": "list of {tag, cells, optional_cells?, note?}; a cell may carry several tags (e.g. switchover_weighted and parallel_product_complement on the same cell). Tests should assert cells ⊆ detected ⊆ cells ∪ optional_cells per tag",
            "untagged_formula_cells": "formula cells that match no catalogue tag (complements, pass-throughs, differences, 1/MTBF); listed so a fingerprinter can be checked for false positives",
            "columns": "role -> column letter; base roles: name, level, count, mtbf_h, mttr_h, fit, lambda_per_h, availability, unavailability, redundancy, switchover_success, auto_fraction, mttr_auto_h, mttr_manual_h, mttr_effective_h, notes; extension roles used here: unit_availability (F2, F3), afr_pct_per_year and availability_pct (F9)",
            "tree_node": {
                "name": "trimmed display name",
                "kind": "leaf | series | parallel | kofn",
                "count": "number of identical copies of the node in series that the engine must apply (A_node = A^count); 1 when the sheet does not apply the quantity",
                "qty_cell": "raw value of the quantity column (informational; equals n on homogeneous redundancy rows, may be ignored by the sheet as in F4)",
                "k, n": "voting parameters; parallel means k = 1; null for leaf and series",
                "homogeneous": "true when the row describes n identical units whose parameters are in params and children is empty",
                "params": "input values exactly as stored (numbers; redundancy text)",
                "derived": "values of formula cells belonging to the row (λ from FIT, MTTR_eff, μ, MTBF from AFR)",
                "cells": "role -> cell address for every parameter and result cell (provenance)",
                "source_row": "1-based row of the node",
                "alt_cells": "other cells computing the same node's availability with another method",
                "children": "ordered list of nodes",
            },
            "evaluation": {
                "leaf": "A_unit = mtbf/(mtbf+mttr), or 1/(1+λ·mttr) when λ is given, or the percent/100; A = A_unit^count",
                "series": "Π A_child",
                "parallel": "homogeneous: 1-(1-A_unit)^n, with switchover success c: c·(1-(1-A_unit)^n) + (1-c)·A_unit; heterogeneous (F6/F7): 1-Π(1-A_child)",
                "kofn": "Σ_{i=k..n} C(n,i)·A_unit^i·(1-A_unit)^(n-i)",
                "ctmc (F5)": "3-state chain with one repair crew; availability = π(2 up)+π(1 up)",
            },
        },
        "method_tags": list(METHOD_TAGS),
        "method_coverage": coverage,
        "workbooks": workbooks,
    }
    (HERE / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"\nmanifest.json: {len(workbooks)} workbooks, "
          f"{sum(len(w['formulas']) for w in workbooks)} formula cells with expected values")
    uncovered = [t for t, ids in coverage.items() if not ids]
    print("method coverage: " + ", ".join(f"{t}:{len(ids)}" for t, ids in coverage.items()))
    assert not uncovered, uncovered
    print("OK")


if __name__ == "__main__":
    main()
